"""Text out of the files people have at hand (§23.4): a PDF datasheet, a Word
document, a spreadsheet, an email, a text file.

Each piece of text keeps where it came from (page, sheet and cell,
paragraph), so a suggestion can point back to it. A file is untrusted:
nothing in it is followed, formulas are read as their cached values or not
at all, and nothing is fetched from links inside it.
"""
from __future__ import annotations

import io
import re
import zipfile
from email import policy
from email.parser import BytesParser
from typing import Optional

MAX_CHARS = 20000
MAX_PAGES = 40
# A file kept whole, as a document's body or the text to rework (extract_document): longer than what the
# assistant reads to suggest fields.
DOCUMENT_CHARS = 60000
DOCUMENT_PAGES = 200
MAX_CELLS = 4000


class UnreadableFile(ValueError):
    pass


def kind_of(filename: str, mime: Optional[str]) -> str:
    name = (filename or "").lower()
    mime = (mime or "").lower()
    if mime.startswith("image/"):
        return "image"
    for ext, kind in ((".pdf", "pdf"), (".docx", "docx"), (".xlsx", "xlsx"), (".xlsm", "xlsx"), (".eml", "email"),
                      (".txt", "text"), (".md", "text"), (".csv", "text"), (".log", "text")):
        if name.endswith(ext):
            return kind
    if mime == "application/pdf":
        return "pdf"
    if mime.startswith("text/"):
        return "text"
    return "unknown"


def extract(content: bytes, filename: str, mime: Optional[str], limit: int = MAX_CHARS) -> tuple[str, list[dict]]:
    """(text, references). The text is marked with its origin, e.g. [page 2]; at most `limit` characters."""
    kind = kind_of(filename, mime)
    if kind == "pdf":
        return _pdf(content, limit)
    if kind == "docx":
        return _docx(content, limit)
    if kind == "xlsx":
        return _xlsx(content, limit)
    if kind == "email":
        return _email(content, limit)
    if kind == "text":
        text = content.decode("utf-8", errors="replace")[:limit]
        return text, [{"kind": "file", "format": "text", "chars": len(text)}]
    raise UnreadableFile("This kind of file cannot be read yet: use a PDF, Word, Excel, email or text file.")


def _page_text(page) -> str:
    """A PDF page's text, line by line as it is laid out. The plain reading follows the file's own text runs,
    and a page written in a word processor with justified lines (Google Docs) comes out a word per line, each
    word its own paragraph; the layout reading keeps the lines, and the spaces it pads columns with are
    folded back to one. The plain reading is the fallback."""
    try:
        laid = page.extract_text(extraction_mode="layout") or ""
    except Exception:
        laid = ""
    if laid.strip():
        lines = [re.sub(r"(?<=\S) {2,}(?=\S)", " ", line.rstrip()) for line in laid.splitlines()]
        margin = min((len(x) - len(x.lstrip()) for x in lines if x.strip()), default=0)
        text = "\n".join(x[margin:] for x in lines)
        return re.sub(r"\n{3,}", "\n\n", text).strip()
    try:
        return (page.extract_text() or "").strip()
    except Exception:
        return ""


def _pdf(content: bytes, limit: int = MAX_CHARS) -> tuple[str, list[dict]]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:  # a damaged or encrypted file is a person-facing message, not a crash
        raise UnreadableFile(f"The PDF could not be opened: {exc}")
    parts, refs, total = [], [], 0
    for n, page in enumerate(reader.pages[:MAX_PAGES if limit <= MAX_CHARS else DOCUMENT_PAGES], start=1):
        text = _page_text(page)
        if not text:
            continue
        chunk = f"[page {n}]\n{text}\n"
        parts.append(chunk)
        refs.append({"kind": "file", "format": "pdf", "page": n, "chars": len(text)})
        total += len(chunk)
        if total >= limit:
            break
    if not parts:
        raise UnreadableFile("The PDF has no text layer (a scan?). Photograph the nameplate or type what it says.")
    return "".join(parts)[:limit], refs


def _docx(content: bytes, limit: int = MAX_CHARS) -> tuple[str, list[dict]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="replace")
    except Exception as exc:
        raise UnreadableFile(f"The Word file could not be opened: {exc}")
    paragraphs = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        text = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", p, re.S))
        text = (text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
                .replace("&quot;", '"').replace("&apos;", "'")).strip()
        if text:
            paragraphs.append(text)
    text = "\n".join(paragraphs)[:limit]
    if not text:
        raise UnreadableFile("The Word file has no text.")
    return text, [{"kind": "file", "format": "docx", "paragraphs": len(paragraphs)}]


def _xlsx(content: bytes, limit: int = MAX_CHARS) -> tuple[str, list[dict]]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)   # cached values, never formulas
    except Exception as exc:
        raise UnreadableFile(f"The spreadsheet could not be opened: {exc}")
    parts, refs, cells = [], [], 0
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue                         # a hidden sheet is not what the person meant to share
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if values:
                rows.append(f"row {r}: " + " | ".join(values))
                cells += len(values)
            if cells >= MAX_CELLS:
                break
        if rows:
            parts.append(f"[sheet {ws.title}]\n" + "\n".join(rows) + "\n")
            refs.append({"kind": "file", "format": "xlsx", "sheet": ws.title, "rows": len(rows)})
        if cells >= MAX_CELLS:
            break
    text = "".join(parts)[:limit]
    if not text:
        raise UnreadableFile("The spreadsheet has no values.")
    return text, refs


def _email(content: bytes, limit: int = MAX_CHARS) -> tuple[str, list[dict]]:
    msg = BytesParser(policy=policy.default).parsebytes(content)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body is not None else ""
    if body is not None and body.get_content_type() == "text/html":
        text = re.sub(r"<[^>]+>", " ", text)
    head = f"Subject: {msg.get('subject', '')}\nDate: {msg.get('date', '')}\n\n"
    text = (head + (text or "")).strip()[:limit]
    return text, [{"kind": "file", "format": "email", "chars": len(text)}]


def extract_document(content: bytes, filename: str, mime: Optional[str]) -> dict:
    """A file as a document's body, in Markdown (intake/markdown.py): its structure kept, as long as a document may
    be. {text, format, truncated}."""
    from app.intake import markdown as md
    name = (filename or "").lower()
    kind = kind_of(filename, mime)
    try:
        if name.endswith((".html", ".htm", ".xhtml")) or (mime or "").lower() == "text/html":
            text, kind = md.html_markdown(content), "html"
        elif name.endswith((".md", ".markdown")):
            text = content.decode("utf-8", errors="replace").replace("\r\n", "\n")
        elif name.endswith(".csv"):
            text = md.csv_markdown(content.decode("utf-8", errors="replace"))
        elif kind == "docx":
            text = md.docx_markdown(content)
        elif kind == "xlsx":
            text = md.xlsx_markdown(content)
        elif kind == "email":
            text = md.email_markdown(content)
        elif kind == "pdf":
            raw, _ = _pdf(content, DOCUMENT_CHARS + 1)
            text = "\n\n".join(md.text_to_markdown(page, wrapped=True)
                                for page in re.split(r"(?m)^\[page \d+\]\n", raw) if page.strip())
        elif kind == "text":
            text = md.text_to_markdown(content.decode("utf-8", errors="replace"))
        else:
            raise UnreadableFile("This kind of file cannot be read yet: use a PDF, Word, Excel, e-mail, web page or "
                                 "text file.")
    except md.Unreadable as exc:
        raise UnreadableFile(str(exc))
    truncated = len(text) > DOCUMENT_CHARS
    return {"text": text[:DOCUMENT_CHARS].strip(), "format": kind, "truncated": truncated}
